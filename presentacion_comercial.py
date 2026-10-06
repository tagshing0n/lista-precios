"""Crea un libro comercial nuevo usando solo las columnas autorizadas."""
import argparse
from decimal import Decimal, InvalidOperation
from pathlib import Path
from math import ceil
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet
import re
import warnings
from io import BytesIO
from openpyxl.drawing.image import Image
from reglas_distribucion import MARCADOR_REDUCIDO, motivos_exclusion

COLUMNAS = ('Nombre de Linea','Codigo de Articulo','Precio de Lista','Marca',
            'Descripcion 1','Descripcion 2','Descripcion 3','Descripcion Extendida',
            'Unidad de Medida','Codigo SAT','Existencia')
ANCHOS = (24,25,17,24,42,42,42,72,15,17,14)
NUMERO = re.compile(r'[+-]?(?:[0-9]+|[0-9]{1,3}(?:,[0-9]{3})+)(?:\.[0-9]+)?')


def disponibilidad(valor):
    if isinstance(valor,bool) or valor is None:
        raise ValueError('Existencia vacia o no numerica; no se puede decidir disponibilidad.')
    texto = str(valor).strip()
    if not NUMERO.fullmatch(texto):
        raise ValueError('Existencia no numerica; no se puede decidir disponibilidad.')
    try:
        numero=Decimal(texto.replace(',',''))
    except InvalidOperation as error:
        raise ValueError('Existencia invalida.') from error
    return '' if numero == 0 else 'D'


def generar_comercial(entrada, salida, logo=None, edicion='Lista general de precios'):
    entrada,salida=Path(entrada),Path(salida)
    if salida.suffix.lower()!='.xlsx' or salida.exists() or entrada.resolve()==salida.resolve():
        raise ValueError('Elige un archivo .xlsx nuevo para la salida comercial.')
    imagen = Image(BytesIO(Path(logo).read_bytes())) if logo else None
    origen=load_workbook(entrada,read_only=True,data_only=False,keep_links=False)
    filas=[]
    try:
        if 'Lista de precios' not in origen.sheetnames:
            raise ValueError('No se encontro la hoja Lista de precios.')
        iterator=origen['Lista de precios'].iter_rows()
        cabecera=[c.value for c in next(iterator)]
        if any(cabecera.count(c)!=1 for c in COLUMNAS):
            raise ValueError('Falta una columna autorizada o hay encabezados duplicados.')
        indices=[cabecera.index(c) for c in COLUMNAS]
        es_disponibilidad = origen.properties.description == MARCADOR_REDUCIDO
        for n,row in enumerate(iterator,2):
            if 'Linea' in cabecera:
                campos_filtro = [row[cabecera.index(c)] for c in ('Linea','Precio de Lista')]
                if any(c.data_type in ('f','e') for c in campos_filtro):
                    raise ValueError(f'Fila {n}: filtro con formula o error.')
                if motivos_exclusion(*(c.value for c in campos_filtro)): continue
            valores=[]
            for nombre,i in zip(COLUMNAS,indices):
                cell=row[i]
                if cell.data_type in ('f','e'):
                    raise ValueError(f'Fila {n}, {nombre}: formula o error requiere revision previa.')
                valor=cell.value
                if nombre=='Existencia':
                    if es_disponibilidad:
                        if valor not in (None,'','D'):
                            raise ValueError(f'Fila {n}: disponibilidad invalida en reducido.')
                        valor=valor or ''
                    else:
                        try:valor=disponibilidad(valor)
                        except ValueError as error:raise ValueError(f'Fila {n}: {error}') from error
                if nombre in ('Codigo de Articulo','Codigo SAT') and valor is not None and not isinstance(valor,str):
                    raise ValueError(f'Fila {n}, {nombre}: el identificador debe venir como texto.')
                valores.append(valor)
            filas.append(valores)
    finally:
        origen.close()
    if len(filas)>1048571:
        raise ValueError('Demasiadas filas para el formato comercial.')
    wb=Workbook(write_only=True)
    wb.properties.creator='listaPrecios'
    wb.properties.title='Lista de precios comercial'
    ws=wb.create_sheet('Lista comercial')
    ws.sheet_view.showGridLines=False
    ws.freeze_panes='C6'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=Worksheet.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1
    ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:5'
    ws.print_options.horizontalCentered=True
    ws.oddFooter.center.text='Pagina &P de &N'
    ws.merged_cells.add('A1:B3')
    ws.merged_cells.add('C1:K2')
    ws.merged_cells.add('C3:K3')
    ws.merged_cells.add('A4:K4')
    if imagen:
        factor=min(280/imagen.width,135/imagen.height)
        imagen.width*=factor;imagen.height*=factor
        ws.add_image(imagen,'A1')
    for i,width in enumerate(ANCHOS,1):
        ws.column_dimensions[chr(64+i)].width=width
    normal=Font(name='Calibri',size=11,color='243746')
    bold=Font(name='Calibri',size=11,bold=True,color='FFFFFF')
    title=Font(name='Calibri',size=22,bold=True,color='FFFFFF')
    left=Alignment(vertical='center',wrap_text=True)
    center=Alignment(horizontal='center',vertical='center',wrap_text=True)
    right=Alignment(horizontal='right',vertical='center')
    dark=PatternFill('solid',fgColor='12264B')
    teal=PatternFill('solid',fgColor='193C63')
    stripe=PatternFill('solid',fgColor='F2F5F9')
    green=Font(name='Calibri',size=11,bold=True,color='176B55')
    def celda(value,font=normal,fill=None,alignment=left):
        c=WriteOnlyCell(ws,value=value)
        if isinstance(value,str):c.data_type='s';c.number_format='@'
        c.font=font;c.alignment=alignment
        if fill:c.fill=fill
        return c
    ws.row_dimensions[1].height=44
    ws.row_dimensions[2].height=40
    ws.row_dimensions[3].height=26
    ws.row_dimensions[4].height=26
    ws.row_dimensions[5].height=34
    ws.append([celda(None),celda(None),celda('ADELAR | LISTA DE PRECIOS',title,dark)]+[celda(None,fill=dark) for _ in range(8)])
    ws.append([celda(None),celda(None)]+[celda(None,fill=dark) for _ in range(9)])
    ws.append([celda(None),celda(None),celda(edicion,bold,PatternFill('solid',fgColor='A62436'))]+[celda(None,fill=PatternFill('solid',fgColor='A62436')) for _ in range(8)])
    ws.append([celda('Disponibilidad: D = disponible | Celda vacia = agotado.  Utiliza los filtros para buscar articulos.')])
    ws.append([celda(c,bold,teal,center) for c in COLUMNAS])
    for n,row in enumerate(filas,6):
        altura=max(27,*[ceil(len(str(row[i] or ''))/(ANCHOS[i]-4))*15+9 for i in (0,1,3,4,5,6,7)])
        ws.row_dimensions[n].height=min(409,altura)
        salida_row=[]
        for i,value in enumerate(row):
            c=celda(value,green if i==10 else normal,stripe if n%2==0 else None,
                    right if i==2 else center if i in (8,9,10) else left)
            if i==2 and isinstance(value,(int,float)):c.number_format='#,##0.00'
            salida_row.append(c)
        ws.append(salida_row)
    if filas:
        ref=f'A5:K{len(filas)+5}'
        t=Table(displayName='CatalogoComercial',ref=ref,
                tableColumns=[TableColumn(id=i,name=c) for i,c in enumerate(COLUMNAS,1)])
        t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=False)
        with warnings.catch_warnings():
            warnings.filterwarnings('ignore',message='In write-only mode you must add table columns manually')
            ws.add_table(t)
        ws.auto_filter.ref=ref
    ws.print_area=f'A1:K{len(filas)+5}'
    creado=False
    try:
        with salida.open('xb') as f:
            creado=True
            wb.save(f)
    except Exception:
        if creado:salida.unlink(missing_ok=True)
        raise
    return {'articulos':len(filas),'disponibles':sum(r[-1]=='D' for r in filas),
            'agotados':sum(r[-1]=='' for r in filas)}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('entrada',type=Path)
    parser.add_argument('--salida',required=True,type=Path)
    parser.add_argument('--logo',type=Path,help='Imagen del logo autorizada por la empresa.')
    parser.add_argument('--edicion',default='Lista general de precios')
    args=parser.parse_args()
    try:
        resumen=generar_comercial(args.entrada,args.salida,args.logo,args.edicion)
    except (ValueError,OSError) as error:
        parser.exit(1,f'Error: {error}\n')
    print('Lista comercial:',args.salida)
    print(resumen)


if __name__=='__main__':main()
