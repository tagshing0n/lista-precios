"""Exportacion reproducible a Excel, conservando identificadores y avisos."""
from pathlib import Path
from math import ceil
from decimal import Decimal
import re
import warnings
import unicodedata
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from validacion_numerica import clasificar

NUMEROS = {'Almacen', 'Linea', 'SubLinea', 'Existencia', 'Precio de Lista'}
DESCRIPCIONES = {'Descripcion 1', 'Descripcion 2', 'Descripcion 3', 'Descripcion Extendida'}
AZUL = '203864'
FUENTE = Font(name='Arial', size=10)
CABECERA = Font(name='Arial', size=10, bold=True, color='FFFFFF')


def texto_seguro(valor, linea, columna):
    if len(valor.encode('utf-16-le')) // 2 > 32767:
        raise ValueError(f'Linea {linea}, {columna}: excede el limite de texto de Excel.')
    if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]', valor):
        raise ValueError(f'Linea {linea}, {columna}: control incompatible con Excel; revisar sin eliminarlo.')
    return valor


def valor_excel(columna, raw, encoding, linea):
    texto = texto_seguro(raw.decode(encoding, errors='strict'), linea, columna)
    if columna not in NUMEROS or not raw:
        return texto, None
    estado, numero, _ = clasificar(columna, raw)
    if estado != 'valido':
        return texto, None
    if columna in {'Almacen', 'Linea', 'SubLinea'} and len(raw) > 1 and raw.startswith(b'0'):
        return texto, 'Identificador con ceros iniciales conservado como texto'
    if len(numero.as_tuple().digits) > 15:
        return texto, 'Numero conservado como texto para evitar perdida de precision en Excel'
    convertido = int(numero) if columna in {'Almacen', 'Linea', 'SubLinea'} else float(numero)
    if Decimal(str(convertido)) != numero:
        return texto, 'Numero conservado como texto para evitar perdida de precision en Excel'
    return convertido, None


def celda(ws, valor, formato=None, encabezado=False, envolver=False):
    c = WriteOnlyCell(ws, value=valor)
    if isinstance(valor, str):
        c.data_type = 's'  # Un texto que comienza con = nunca se convierte en formula.
        c.number_format = '@'
    if formato and not isinstance(valor, str):
        c.number_format = formato
    c.font = CABECERA if encabezado else FUENTE
    c.alignment = Alignment(horizontal='center' if encabezado else ('left' if isinstance(valor, str) else 'right'),
                            vertical='center', wrap_text=envolver or encabezado)
    if encabezado:
        c.fill = PatternFill('solid', fgColor=AZUL)
    return c


def tabla(ws, nombre, encabezados, fin, fila_encabezado=1):
    ref = f'A{fila_encabezado}:{get_column_letter(len(encabezados))}{fin}'
    t = Table(displayName=nombre, ref=ref,
              tableColumns=[TableColumn(id=i, name=c) for i,c in enumerate(encabezados,1)])
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', message='In write-only mode you must add table columns manually')
        ws.add_table(t)
    ws.auto_filter.ref = ref


def clave_orden(row):
    """Numeros primero; textos invalidos y vacios al final, sin descartarlos."""
    def numero(columna, valor):
        raw = str(valor).encode('ascii', errors='replace')
        estado, valor_decimal, _ = clasificar(columna, raw)
        if estado == 'valido':
            return (0, valor_decimal, '')
        return (2 if valor == '' else 1, Decimal(0), str(valor).casefold())
    return (numero('Linea', row[1]), numero('SubLinea', row[3]),
            unicodedata.normalize('NFC', str(row[6])).casefold())


def exportar_excel(entrada, destino, encabezados, numerico, texto, encoding='cp1252', fuente=None,
                   reducida=False):
    if destino.exists() or destino.resolve() == entrada.resolve():
        raise ValueError('El Excel de salida ya existe o coincide con la entrada.')
    # Validacion completa antes de iniciar la escritura del libro.
    rows = []
    extra = []
    for n,linea in enumerate(entrada.read_bytes().split(b'\n'),1):
        if not linea and n > numerico['filas'] + 1:
            continue
        campos = linea.removesuffix(b'\r').split(b'|')
        if len(campos) != len(encabezados):
            raise ValueError(f'Linea {n}: estructura incompatible con Excel.')
        if n == 1:
            if campos != [c.encode('ascii') for c in encabezados]:
                raise ValueError('Los encabezados no coinciden.')
            continue
        valores = []
        for col,raw in zip(encabezados,campos):
            valor, aviso = valor_excel(col, raw, encoding, n)
            valores.append(valor)
            if aviso:
                extra.append((n, col, aviso, valor, ''))
        rows.append(valores)
    if len(rows) != numerico['filas'] or len(rows) > 1048575:
        raise ValueError('Cantidad de filas incompatible con el reporte o con Excel.')
    issues = []
    for i in numerico['incidencias']:
        motivo = ', '.join(i['avisos']) if i['avisos'] else 'Valor ' + i['estado']
        raw = bytes.fromhex(i['bytes_originales_hex'])
        issues.append((i['linea'], i['columna'], motivo, raw.decode(encoding), ''))
    for i in texto['incidencias']:
        raw = bytes.fromhex(i['valor_original_bytes_hex'])
        issues.append((i['linea'], i['columna'],
                       'Revisar lectura de caracteres no ASCII (no implica error)',
                       raw.decode(encoding), raw.decode('cp850')))
    for n,row in enumerate(rows,2):
        if not row[2]:
            issues.append((n, 'Nombre de Linea', 'Nombre de linea vacio: revisar', '', ''))
    issues.extend(extra)
    for n,col,motivo,valor,alternativa in issues:
        texto_seguro(str(valor), n, col)
        texto_seguro(alternativa, n, col)
    if len(issues) > 1048571:
        raise ValueError('Las incidencias exceden la capacidad de la hoja.')
    excluidas = {'Existencia'}
    if reducida:
        excluidas.update({'Nombre de SubLinea', 'Codigo Fabricante'})
    columnas_salida = [c for c in encabezados if c not in excluidas] + ['Existencia']
    indices_salida = [encabezados.index(c) for c in columnas_salida]
    ordenadas = sorted(enumerate(rows, 2), key=lambda item: clave_orden(item[1]))
    filas_excel = {linea: fila for fila, (linea, _) in enumerate(ordenadas, 2)}
    wb = Workbook(write_only=True)
    wb.properties.title = 'Lista de precios'
    wb.properties.creator = 'listaPrecios'
    datos = wb.create_sheet('Lista de precios')
    revision = wb.create_sheet('Incidencias')
    datos.sheet_view.showGridLines = False
    revision.sheet_view.showGridLines = False
    datos.sheet_properties.tabColor = AZUL
    revision.sheet_properties.tabColor = 'D99B32'
    datos.freeze_panes = f'{get_column_letter(columnas_salida.index("Codigo de Articulo") + 2)}2'
    revision.freeze_panes = 'D6'
    widths = [11,11,25,13,28,16,23,19,24,43,43,43,76,18,18,25]
    widths = [widths[i] for i in indices_salida]
    for i,ancho in enumerate(widths,1):
        datos.column_dimensions[get_column_letter(i)].width = ancho
    datos.row_dimensions[1].height = 32
    datos.append([celda(datos,c,encabezado=True) for c in columnas_salida])
    for n,(_, original_row) in enumerate(ordenadas,2):
        row = [original_row[i] for i in indices_salida]
        altura = max([24] + [ceil(len(str(row[i])) / (widths[i] - 3)) * 13 + 8 for i,c in enumerate(columnas_salida) if c in DESCRIPCIONES])
        datos.row_dimensions[n].height = min(altura, 400)
        celdas = []
        for i,v in enumerate(row):
            formato = '#,##0.000' if columnas_salida[i] == 'Existencia' else '#,##0.00' if columnas_salida[i] == 'Precio de Lista' else '0'
            celdas.append(celda(datos,v,formato,envolver=columnas_salida[i] in DESCRIPCIONES))
        datos.append(celdas)
    if rows:
        tabla(datos,'ListaPrecios',columnas_salida,len(rows)+1)
        precio = get_column_letter(columnas_salida.index('Precio de Lista') + 1)
        datos.conditional_formatting.add(f'{precio}2:{precio}{len(rows)+1}', FormulaRule(
            formula=[f'AND(ISNUMBER({precio}2),{precio}2=0)'], fill=PatternFill('solid',fgColor='FFF2CC')))
    headers = ['Linea TXT','Fila Excel','Codigo de Articulo','Columna','Motivo','Valor mostrado','Vista alternativa CP850','Revision']
    issue_widths = [12,12,24,27,52,72,72,18]
    for i,ancho in enumerate(issue_widths,1):
        revision.column_dimensions[get_column_letter(i)].width = ancho
    for n in range(1,5):
        revision.row_dimensions[n].height = 24
    revision.append([celda(revision,'Incidencias pendientes')])
    revision.append([celda(revision,'Fuente: ' + (fuente or entrada.name))])
    revision.append([celda(revision,f'Lectura {encoding} provisional. Comparar caracteres dudosos; no se aplican nuevas correcciones.')])
    revision.append([celda(revision,f'Filas: {len(rows)}. Observaciones: {len(issues)}. La revision se regenera en cada ejecucion.')])
    revision.row_dimensions[5].height = 32
    revision.append([celda(revision,c,encabezado=True) for c in headers])
    for fila,(n,col,motivo,valor,alternativa) in enumerate(issues,6):
        codigo = str(rows[n-2][6])
        valores = [n,filas_excel[n],codigo,col,motivo,valor,alternativa,'Pendiente']
        revision.row_dimensions[fila].height = min(400,max(28,*[
            ceil(len(str(valores[i]))/(issue_widths[i]-3))*13+8 for i in (4,5,6)]))
        revision.append([celda(revision,v,envolver=i in (4,5,6)) for i,v in enumerate(valores)])
    if issues:
        tabla(revision,'IncidenciasPrecios',headers,len(issues)+5,5)
        dv = DataValidation(type='list', formula1='"Pendiente,Revisado"', allow_blank=False)
        revision.data_validations.append(dv)
        dv.add(f'H6:H{len(issues)+5}')
    # No formulas de datos; las reglas de formato las calcula Excel al abrirlo.
    creado = False
    try:
        with destino.open('xb') as archivo:
            creado = True
            wb.save(archivo)
    except Exception:
        if creado:
            destino.unlink(missing_ok=True)
        raise
    return {'filas': len(rows), 'incidencias': len(issues), 'lectura_provisional': encoding}
