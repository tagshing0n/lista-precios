from pathlib import Path
from datetime import datetime
import uuid
import shutil
from contextlib import contextmanager
import unittest
from unittest.mock import patch
import json
from almacenes import seleccionar_almacen
from PIL import Image
from openpyxl import load_workbook
import listaPrecios
from generar_listas import ejecutar

@contextmanager
def temporal(work):
    base=work/('test_ejecutor_'+uuid.uuid4().hex)
    base.mkdir()
    try: yield base
    finally:
        assert base.resolve().is_relative_to(work.resolve())
        assert base.name.startswith('test_ejecutor_')
        shutil.rmtree(base)


class EjecutorTests(unittest.TestCase):
    def test_generacion_repetible_y_archivos_faltantes(self):
        work = Path(__file__).resolve().parents[1] / 'work'
        work.mkdir(exist_ok=True)
        with temporal(work) as folder:
            base=Path(folder)
            with self.assertRaises(ValueError): ejecutar(base)
            self.assertFalse((base/'outputs').exists())
            datos=base/'datos';datos.mkdir()
            Image.new('RGB',(100,50),'white').save(datos/'logo_empresa.jpeg')
            fila=['1','2','LINEA','201','SUB','-1','00001','12.50','MARCA','ARTICULO','','','DETALLE','PZ','01010101','FAB']
            original=('|'.join(listaPrecios.ENCABEZADOS)+'\n'+'|'.join(fila)).encode()
            (datos/'Lista-Precios.txt').write_bytes(original)
            primera=ejecutar(base,datetime(2026,10,6,10))
            segunda=ejecutar(base,datetime(2026,10,6,11),almacen='egarama-queretaro')
            estado=json.loads((segunda/'ejecucion.json').read_text(encoding='utf-8'))
            self.assertEqual(estado['almacen'],'egarama-queretaro')
            otro=load_workbook(segunda/'Lista_Precios_Final_EGARAMA_QUERÉTARO_2026-10-06.xlsx')
            try:
                self.assertEqual(otro.active['C1'].value,'EGARAMA QUERÉTARO | LISTA DE PRECIOS')
            finally:otro.close()
            self.assertNotEqual(primera,segunda)
            with self.assertRaises(FileExistsError): ejecutar(base,datetime(2026,10,6,10))
            self.assertEqual((datos/'Lista-Precios.txt').read_bytes(),original)
            for nombre in ('Lista_Precios_ADELAR_2026-10-06.xlsx','Lista_Precios_reducida_ADELAR_2026-10-06.xlsx','Lista_Precios_Final_ADELAR_2026-10-06.xlsx'):
                self.assertTrue((primera/nombre).is_file())
            manifiesto=json.loads((primera/'Lista_Precios_ADELAR_2026-10-06_auditoria'/'proceso.json').read_text(encoding='utf-8'))
            self.assertEqual(manifiesto['excel_reducido'],'Lista_Precios_reducida_ADELAR_2026-10-06.xlsx')
            for nombre in estado['archivos'].values():
                self.assertTrue((segunda/nombre).is_file())
            wb=load_workbook(primera/'Lista_Precios_Final_ADELAR_2026-10-06.xlsx')
            try:
                self.assertEqual(wb.active['K6'].value,'D')
                self.assertEqual(wb.active['B6'].value,'00001')
                self.assertEqual(wb.active['J6'].value,'01010101')
                self.assertEqual(wb.active['C6'].font.name,'Open Sans Light')
                self.assertIn('Los precios de lista no incluyen IVA',wb.active['A4'].value)
            finally:wb.close()

    def test_menu_reintenta_y_permite_salir(self):
        with patch('builtins.input',side_effect=['', '4', '2']):
            self.assertEqual(seleccionar_almacen(),'egarama-mexico')
        with patch('builtins.input',return_value='0'):
            self.assertIsNone(seleccionar_almacen())
