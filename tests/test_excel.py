from pathlib import Path
import contextlib
import io
import shutil
import unittest
import uuid
import json
from revisiones_aceptadas import identidad
from openpyxl import load_workbook
import listaPrecios
from proceso_completo import generar_excel
from exportacion_excel import valor_excel, texto_seguro


class ValoresTests(unittest.TestCase):
    def test_identificadores_y_numeros(self):
        self.assertEqual(valor_excel('Codigo SAT',b'00123456','cp1252',2)[0],'00123456')
        self.assertEqual(valor_excel('Precio de Lista',b'1,234.50','cp1252',2)[0],1234.5)
        self.assertEqual(valor_excel('Precio de Lista',b'INVALIDO','cp1252',2)[0],'INVALIDO')
        self.assertEqual(valor_excel('Almacen',b'001','cp1252',2)[0],'001')

    def test_precision_y_texto_demasiado_largo(self):
        valor,aviso = valor_excel('Existencia',b'1234567890123456','cp1252',2)
        self.assertEqual(valor,'1234567890123456')
        self.assertTrue(aviso)
        with self.assertRaises(ValueError):
            texto_seguro('a' * 32768, 2, 'Descripcion 1')

    def test_control_no_aprobado_no_se_borra(self):
        with self.assertRaises(ValueError):
            texto_seguro('A\x00B',2,'Descripcion 1')


class ProcesoTests(unittest.TestCase):
    def test_aceptaciones_persistentes_solo_para_casos_revisados(self):
        self.campos[2] = b''
        conocidos = [c.decode('cp1252').strip(' ') for c in self.campos]
        nuevos = self.campos.copy()
        nuevos[6] = b'NUEVO'
        self.entrada.write_bytes(self.cabecera+b'|'.join(nuevos)+b'\n'+b'|'.join(self.campos))
        reglas={k:[identidad(conocidos)] for k in ('precio_cero','nombre_linea_vacio')}
        self.entrada.with_suffix('.revisiones.json').write_text(json.dumps({'version':1,'aceptaciones':reglas}))
        resultado=self.ejecutar()
        self.assertEqual(resultado['aceptaciones_aplicadas'],{'precio_cero':1,'nombre_linea_vacio':1})
        wb=load_workbook(self.destino)
        self.addCleanup(wb.close)
        pendientes=list(wb['Incidencias'].iter_rows(min_row=6,values_only=True))
        self.assertEqual(len(pendientes),2)
        self.assertTrue(all(r[2]=='NUEVO' for r in pendientes))
        self.assertEqual(wb['Lista de precios']['G2'].value,0)
        self.assertIsNone(wb['Lista de precios']['C2'].value)
        self.assertEqual(len(wb['Lista de precios'].conditional_formatting),1)

    def setUp(self):
        self.work = Path(__file__).resolve().parents[1] / 'work'
        self.root = self.work / ('test_excel_' + uuid.uuid4().hex)
        self.root.mkdir(parents=True)
        self.addCleanup(self.limpiar)
        self.entrada = self.root / 'entrada.txt'
        self.destino = self.root / 'salida.xlsx'
        self.campos = [b'1',b'2',b' LINEA ',b'201',b'SUBLINEA',b'1,234.000',
                       b'000123',b'0.00',b'MARCA',b'PI\xa5ON\tDEMO',b'',b'',b'.',
                       b'PZ',b'00123456',b'=CODIGO']
        self.cabecera = b'|'.join(c.encode() for c in listaPrecios.ENCABEZADOS) + b'\r\n'

    def limpiar(self):
        root = self.root.resolve()
        if not root.is_relative_to(self.work.resolve()) or not root.name.startswith('test_excel_'):
            raise AssertionError('Ruta temporal no valida')
        shutil.rmtree(root)

    def ejecutar(self):
        with contextlib.redirect_stdout(io.StringIO()):
            return generar_excel(self.entrada,self.destino,listaPrecios.ENCABEZADOS,
                                 listaPrecios.COLUMNAS_NUMERICAS,listaPrecios.COLUMNAS_TEXTO,
                                 listaPrecios.corregir_estructura)

    def test_pipeline_tipos_columnas_y_formula_literal(self):
        campos = self.campos[:11] + [b''] + self.campos[11:]
        original = self.cabecera + b'|'.join(campos)
        self.entrada.write_bytes(original)
        resultado = self.ejecutar()
        self.assertEqual(self.entrada.read_bytes(),original)
        wb = load_workbook(self.destino)
        self.addCleanup(wb.close)
        ws = wb['Lista de precios']
        heads = [c.value for c in ws[1]]
        self.assertEqual(len(heads),16)
        self.assertEqual(heads[4],'Nombre de SubLinea')
        self.assertEqual(heads[-1],'Existencia')
        values = {head: ws.cell(2,i+1) for i,head in enumerate(heads)}
        self.assertEqual(values['Codigo de Articulo'].value,'000123')
        self.assertEqual(values['Codigo SAT'].value,'00123456')
        self.assertEqual(values['Codigo Fabricante'].value,'=CODIGO')
        self.assertEqual(values['Codigo Fabricante'].data_type,'s')
        self.assertEqual(values['Precio de Lista'].value,0)
        self.assertEqual(values['Precio de Lista'].data_type,'n')
        self.assertEqual(values['Existencia'].value,1234)
        self.assertEqual(values['Descripcion 1'].value,'PI\u00d1ON DEMO')
        self.assertGreater(resultado['incidencias'],0)
        self.assertEqual(ws.tables['ListaPrecios'].ref,'A1:P2')
        reducido = load_workbook(self.root/'salida_reducida.xlsx')
        self.addCleanup(reducido.close)
        headers_reducidos = [c.value for c in reducido['Lista de precios'][1]]
        self.assertEqual(len(headers_reducidos),14)
        self.assertNotIn('Nombre de SubLinea',headers_reducidos)
        self.assertNotIn('Codigo Fabricante',headers_reducidos)
        self.assertEqual(headers_reducidos[-1],'Existencia')

    def test_orden_multinivel_y_referencias_de_incidencias(self):
        originales = []
        for linea,sublinea,codigo in [('10','1','A'),('2','10','A'),('2','2','B'),
                                      ('02','2','a'),('2','2','A'),('2','2','001')]:
            campos = self.campos.copy()
            campos[1],campos[3],campos[6] = linea.encode(),sublinea.encode(),codigo.encode()
            originales.append(campos)
        self.entrada.write_bytes(self.cabecera+b'\n'.join(b'|'.join(r) for r in originales))
        self.ejecutar()
        for nombre in ('salida.xlsx','salida_reducida.xlsx'):
            wb = load_workbook(self.root/nombre)
            self.addCleanup(wb.close)
            ws = wb['Lista de precios']
            heads = [c.value for c in ws[1]]
            rows = list(ws.iter_rows(min_row=2,values_only=True))
            codigos = [r[heads.index('Codigo de Articulo')] for r in rows]
            self.assertEqual(codigos,['001','a','A','B','A','A'])
            self.assertEqual([r[1] for r in rows],[2,'02',2,2,2,10])
            for issue in wb['Incidencias'].iter_rows(min_row=6,values_only=True):
                origen,fila,codigo = issue[:3]
                self.assertEqual(originales[origen-2][6].decode(),codigo)
                self.assertEqual(ws.cell(fila,heads.index('Codigo de Articulo')+1).value,codigo)
                self.assertEqual(str(ws.cell(fila,2).value), originales[origen-2][1].decode())

    def test_rechaza_reducido_existente_antes_de_crear_archivos(self):
        self.entrada.write_bytes(self.cabecera+b'|'.join(self.campos))
        reducido = self.root/'salida_reducida.xlsx'
        reducido.write_bytes(b'CONSERVAR')
        with self.assertRaises(ValueError):
            self.ejecutar()
        self.assertEqual(reducido.read_bytes(),b'CONSERVAR')
        self.assertFalse(self.destino.exists())
        self.assertFalse((self.root/'salida_auditoria').exists())

    def test_fallo_estructural_no_genera_excel(self):
        campos = self.campos[:10] + [b''] + self.campos[10:]
        campos[11:14] = [b'SEGUNDA',b'TERCERA',b'']
        self.entrada.write_bytes(self.cabecera + b'|'.join(campos))
        with self.assertRaises(ValueError):
            self.ejecutar()
        self.assertFalse(self.destino.exists())
        self.assertTrue((self.root / 'salida_auditoria' / '01_estructura.incidencias.json').exists())

    def test_rechaza_destino_existente(self):
        self.entrada.write_bytes(self.cabecera+b'|'.join(self.campos))
        self.destino.write_bytes(b'CONSERVAR')
        with self.assertRaises(ValueError):
            self.ejecutar()
        self.assertEqual(self.destino.read_bytes(),b'CONSERVAR')

    def test_excel_omite_validos_pero_mantiene_campo_mixto(self):
        self.campos[7] = b'10.00'
        self.campos[9] = 'CONEXIÓN CIGÜEÑAL JIC 37ø'.encode('cp1252')
        self.entrada.write_bytes(self.cabecera+b'|'.join(self.campos))
        self.assertEqual(self.ejecutar()['incidencias'],0)
        wb = load_workbook(self.destino)
        self.addCleanup(wb.close)
        self.assertEqual(wb['Lista de precios']['I2'].value,'CONEXIÓN CIGÜEÑAL JIC 37°')
        self.campos[9] += ' REPøFNO'.encode('cp1252')
        self.entrada.write_bytes(self.cabecera+b'|'.join(self.campos))
        self.destino = self.root/'mixto.xlsx'
        self.assertEqual(self.ejecutar()['incidencias'],1)
