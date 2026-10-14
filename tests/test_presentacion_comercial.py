from pathlib import Path
import unittest,uuid,shutil,hashlib
from zipfile import ZipFile
from openpyxl import Workbook,load_workbook
from presentacion_comercial import COLUMNAS,generar_comercial,disponibilidad
from reglas_distribucion import MARCADOR_REDUCIDO

class ComercialTests(unittest.TestCase):
    def test_almacenes_conservan_datos_y_aplican_colores(self):
        self.fuente()
        datos=None
        for clave,nombre,titulo,encabezado,edicion in (
            ('adelar','ADELAR','12264B','193C63','A62436'),
            ('egarama-mexico','EGARAMA MÉXICO','008000','008000','A62436'),
            ('egarama-queretaro','EGARAMA QUERÉTARO','ED2100','ED2100','12264B'),
        ):
            destino=self.root/(clave+'.xlsx')
            generar_comercial(self.source,destino,almacen=clave)
            wb=load_workbook(destino)
            try:
                ws=wb.active
                self.assertEqual(ws['C1'].value,nombre+' | LISTA DE PRECIOS')
                for celda,color in (('C1',titulo),('A5',encabezado),('C3',edicion)):
                    self.assertEqual(ws[celda].fill.fgColor.rgb[-6:],color)
                self.assertEqual(ws['A4'].value,'Disponibilidad: D = disponible | Celda vacia = agotado. Los precios de lista no incluyen IVA')
                actuales=list(ws.iter_rows(min_row=5,values_only=True))
                if datos is None:datos=actuales
                self.assertEqual(actuales,datos)
                self.assertIsNone(ws.tables['CatalogoComercial'].tableStyleInfo.name)
            finally:wb.close()

    def setUp(self):
        self.work=Path(__file__).resolve().parents[1]/'work'
        self.root=self.work/('test_comercial_'+uuid.uuid4().hex)
        self.root.mkdir(parents=True)
        self.addCleanup(self.limpiar)
        self.source=self.root/'origen.xlsx';self.dest=self.root/'comercial.xlsx'

    def limpiar(self):
        assert self.root.resolve().is_relative_to(self.work.resolve())
        assert self.root.name.startswith('test_comercial_')
        shutil.rmtree(self.root)

    def fuente(self,valor=0):
        wb=Workbook();ws=wb.active;ws.title='Lista de precios'
        ws.append(['Almacen',*COLUMNAS,'Codigo Fabricante'])
        for i,cantidad in enumerate((valor,4,-1,0.5)):
            ws.append(['SECRETO-ALMACEN','LINEA','000'+str(i),12.35,'MARCA',
                       '=TEXTO','DESC 2','DESC 3','DETALLE','PZ','00123456',cantidad,'SECRETO-FABRICANTE'])
            ws.cell(i+2,6).data_type='s'
        wb.create_sheet('Incidencias')['A1']='SECRETO-OBSERVACION'
        wb.save(self.source)

    def test_disponibilidad_y_errores_no_se_confunden_con_agotados(self):
        for valor in (0,'0.000','-0.0'):self.assertEqual(disponibilidad(valor),'')
        for valor in (-2,0.5,'1,000.00'):self.assertEqual(disponibilidad(valor),'D')
        for valor in (None,True,'','NaN','1,23','=1+1'):
            with self.assertRaises(ValueError):disponibilidad(valor)

    def test_solo_columnas_autorizadas_y_cantidades_no_exportadas(self):
        self.fuente()
        original=hashlib.sha256(self.source.read_bytes()).hexdigest()
        self.assertEqual(generar_comercial(self.source,self.dest),{'articulos':4,'disponibles':3,'agotados':1})
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(),original)
        wb=load_workbook(self.dest);self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames,['Lista comercial'])
        ws=wb.active
        self.assertEqual(tuple(c.value for c in ws[5]),COLUMNAS)
        self.assertEqual([ws.cell(i,11).value for i in range(6,10)],[None,'D','D','D'])
        self.assertEqual(ws['B6'].value,'0000')
        self.assertEqual(ws['J6'].value,'00123456')
        self.assertEqual(ws['E6'].value,'=TEXTO')
        self.assertEqual(ws['E6'].data_type,'s')
        self.assertEqual(ws['E5'].value,'Descripcion 1')
        self.assertEqual(ws['F6'].value,'DESC 2')
        self.assertEqual(ws['C6'].value,12.35)
        self.assertEqual(ws.tables['CatalogoComercial'].ref,'A5:K9')
        with ZipFile(self.dest) as z:
            self.assertFalse(any('externalLink' in n or 'comments' in n for n in z.namelist()))
            self.assertNotIn(b'SECRETO',b''.join(z.read(n) for n in z.namelist()))

    def test_no_sobrescribe_ni_genera_salida_con_existencia_invalida(self):
        self.fuente('DESCONOCIDO')
        with self.assertRaises(ValueError):generar_comercial(self.source,self.dest)
        self.assertFalse(self.dest.exists())
        self.dest.write_bytes(b'CONSERVAR')
        with self.assertRaises(ValueError):generar_comercial(self.source,self.dest)
        self.assertEqual(self.dest.read_bytes(),b'CONSERVAR')

    def test_reducido_disponibilidad_logo_y_filtros(self):
        from PIL import Image
        self.fuente()
        wb=load_workbook(self.source)
        wb.properties.description=MARCADOR_REDUCIDO
        ws=wb.active
        ws.cell(1,14,'Linea')
        for n,(linea,precio,stock) in enumerate([(2,12,'D'),(2,12,None),(85,12,'D'),(2,1,'D')],2):
            ws.cell(n,14,linea);ws.cell(n,4,precio)
            ws.cell(n,12).value=stock
        wb.save(self.source);wb.close()
        logo=self.root/'logo.png';Image.new('RGB',(200,100),'white').save(logo)
        result=generar_comercial(self.source,self.dest,logo,'Primera edicion')
        self.assertEqual(result,{'articulos':2,'disponibles':1,'agotados':1})
        wb=load_workbook(self.dest);self.addCleanup(wb.close)
        self.assertEqual(len(wb.active._images),1)
        self.assertEqual(wb.active['C3'].value,'Primera edicion')
